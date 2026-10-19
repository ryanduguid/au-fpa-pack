import json
from pathlib import Path
from unittest import mock

import pytest

ROOT = Path(__file__).resolve().parents[1]


def output_json(result) -> dict:
    assert result.stdout, result.stderr
    return json.loads(result.stdout)


@pytest.mark.parametrize("aliased", [False, True])
@pytest.mark.parametrize("existing", [False, True])
def test_model_export_writes_xlsx_with_both_sheets(tmp_path, run_cli, existing, aliased):
    from openpyxl import load_workbook

    assert run_cli("init", str(tmp_path)).returncode == 0
    config_path = ROOT / "examples/ridgeline/config.yaml"
    if aliased:
        source = config_path.read_text(encoding="utf-8").replace(
            "name: Ridgeline Chair Co.",
            "name: &field !!str\n  ? &selector !!value {=: tax_rate}\n"
            "  : growth_rate\n  =: cogs_pct",
        ).replace("tax_rate: 0.21", "? *selector\n: 0.21").replace(
            "    cogs_pct: ", "    ? *field\n    : ", 1,
        )
        config_path = tmp_path / "valid-alias.yaml"
        config_path.write_text(source, encoding="utf-8")
    before = config_path.read_bytes()
    out_path = tmp_path / "model.xlsx"
    if existing:
        out_path.write_bytes(b"previous report")

    result = run_cli(
        "model-export",
        str(tmp_path),
        "--config",
        str(config_path),
        "--out",
        str(out_path),
    )

    assert result.returncode == 0
    payload = output_json(result)
    assert payload["ok"] is True
    assert payload["command"] == "model-export"
    data = payload["data"]
    assert set(data["sheets"]) == {"Assumptions", "Model"}
    assert out_path.exists()
    wb = load_workbook(out_path)
    assert set(wb.sheetnames) == {"Assumptions", "Model"}
    wb.close()
    assert config_path.read_bytes() == before


def test_model_export_fails_without_workspace(tmp_path, run_cli):
    config_path = ROOT / "examples/ridgeline/config.yaml"
    out_path = tmp_path / "model.xlsx"

    result = run_cli(
        "model-export",
        str(tmp_path),
        "--config",
        str(config_path),
        "--out",
        str(out_path),
    )

    assert result.returncode == 1
    payload = output_json(result)
    assert payload["ok"] is False
    assert payload["error"]["type"] == "workspace_not_initialized"


def test_model_export_fails_on_invalid_config(tmp_path, run_cli):
    assert run_cli("init", str(tmp_path)).returncode == 0
    bad_config = tmp_path / "bad.yaml"
    bad_config.write_text("not: valid: entity: config\n")
    out_path = tmp_path / "model.xlsx"

    result = run_cli(
        "model-export",
        str(tmp_path),
        "--config",
        str(bad_config),
        "--out",
        str(out_path),
    )

    assert result.returncode == 1
    payload = output_json(result)
    assert payload["ok"] is False


@pytest.mark.parametrize("alias_kind", ["same", "relative", "source symlink", "output symlink", "hard link"])
def test_model_export_preserves_configuration_when_output_aliases_it(tmp_path, run_cli, alias_kind):
    assert run_cli("init", str(tmp_path)).returncode == 0
    config = tmp_path / "config.yaml"
    original = (ROOT / "examples/ridgeline/config.yaml").read_bytes()
    config.write_bytes(original)
    source, output = config, config
    try:
        if alias_kind == "relative":
            (tmp_path / "child").mkdir()
            output = tmp_path / "child" / ".." / config.name
        elif alias_kind == "source symlink":
            source = tmp_path / "config-link.yaml"
            source.symlink_to(config)
        elif alias_kind == "output symlink":
            output = tmp_path / "model.xlsx"
            output.symlink_to(config)
        elif alias_kind == "hard link":
            output = tmp_path / "model.xlsx"
            output.hardlink_to(config)
    except OSError as error:
        pytest.skip(f"Cannot create a local file link: {error}")

    result = run_cli("model-export", str(tmp_path), "--config", str(source), "--out", str(output))

    assert config.read_bytes() == original
    assert source.read_bytes() == original
    assert result.returncode == 1
    payload = output_json(result)
    assert payload["ok"] is False
    assert payload["error"]["type"] == "export_failed"


@pytest.mark.parametrize("existing", [False, True])
@pytest.mark.parametrize("failure,diagnostic", [
    ("\ntax_rate: 0.3\n", "tax_rate"),
    ("\n? !!int bad\n: ignored\n", "invalid literal"),
    ("\n? [key]\n: ignored\n", "unhashable key"),
    ("\n? !!python/object:builtins.object {}\n: ignored\n", "could not determine"),
])
def test_invalid_yaml_config_fails_before_model_export_and_preserves_files(
    tmp_path, run_cli, monkeypatch, existing, failure, diagnostic,
):
    assert run_cli("init", str(tmp_path)).returncode == 0
    config = tmp_path / "duplicate.yaml"
    config.write_bytes((ROOT / "examples/ridgeline/config.yaml").read_bytes() + failure.encode("utf-8"))
    original = config.read_bytes()
    output = tmp_path / "model.xlsx"
    if existing:
        output.write_bytes(b"previous report")
    export = mock.Mock(side_effect=AssertionError("Invalid input must not reach workbook generation"))
    monkeypatch.setattr("pyfpa.excel.model_workbook.model_to_excel", export)

    result = run_cli("model-export", str(tmp_path), "--config", str(config), "--out", str(output))

    assert result.returncode == 1
    payload = output_json(result)
    assert payload["command"] == "model-export"
    assert payload["ok"] is False
    assert payload["error"]["type"] == "invalid_config"
    assert diagnostic in payload["error"]["message"]
    export.assert_not_called()
    assert config.read_bytes() == original
    if existing:
        assert output.read_bytes() == b"previous report"
    else:
        assert not output.exists()


@pytest.mark.parametrize("existing", [False, True])
def test_value_tag_alias_duplicate_fails_before_export_and_preserves_files(tmp_path, run_cli, monkeypatch, existing):
    assert run_cli("init", str(tmp_path)).returncode == 0
    source = (ROOT / "examples/ridgeline/config.yaml").read_text(encoding="utf-8")
    assert "name: Ridgeline Chair Co." in source
    assert "tax_rate: 0.21" in source and "    growth_rate: 0.10" in source
    source = source.replace(
        "name: Ridgeline Chair Co.",
        "name: &field !!str\n  ? &tax_key !!value {=: tax_rate}\n"
        "  : annotation_only\n  =: growth_rate",
    ).replace("tax_rate: 0.21", "? *tax_key\n: 0.21").replace(
        "    growth_rate: 0.10", "    ? *field\n    : 0.10\n    growth_rate: 0.0", 1,
    )
    config = tmp_path / "duplicate.yaml"
    config.write_text(source, encoding="utf-8")
    before = config.read_bytes()
    output = tmp_path / "model.xlsx"
    if existing:
        output.write_bytes(b"previous report")
    export = mock.Mock(side_effect=AssertionError("Duplicate input must not reach workbook generation"))
    monkeypatch.setattr("pyfpa.excel.model_workbook.model_to_excel", export)

    from pyfpa.config.schemas import EntityConfig

    with mock.patch.object(EntityConfig, "model_validate", wraps=EntityConfig.model_validate) as validate:
        result = run_cli("model-export", str(tmp_path), "--config", str(config), "--out", str(output))
    try:
        assert result.returncode == 1
        payload = output_json(result)
        assert payload["error"]["type"] == "invalid_config"
        assert "duplicate" in payload["error"]["message"] and "growth_rate" in payload["error"]["message"]
        export.assert_not_called()
        validate.assert_not_called()
    finally:
        assert config.read_bytes() == before
        if existing:
            assert output.read_bytes() == b"previous report"
        else:
            assert not output.exists()


def test_failed_workbook_inspection_keeps_previous_export(tmp_path, run_cli, monkeypatch):
    import pyfpa.excel.model_workbook as reporting

    assert run_cli("init", str(tmp_path)).returncode == 0
    out_path = tmp_path / "model.xlsx"
    out_path.write_bytes(b"previous report")
    before = set(tmp_path.iterdir())

    def fail_read(path, **kwargs):
        raise OSError("cannot inspect workbook")

    monkeypatch.setattr(reporting, "load_workbook", fail_read)
    result = run_cli("model-export", str(tmp_path), "--config",
                     str(ROOT / "examples/ridgeline/config.yaml"), "--out", str(out_path))
    assert result.returncode == 1
    assert output_json(result)["error"]["type"] == "export_failed"
    assert out_path.read_bytes() == b"previous report"
    assert set(tmp_path.iterdir()) == before
