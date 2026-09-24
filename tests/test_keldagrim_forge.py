from pathlib import Path

import pytest

pytest.importorskip("formulas")

import json  # noqa: E402
import subprocess  # noqa: E402
import sys  # noqa: E402

import yaml  # noqa: E402

from pyfpa.config.loader import load_config  # noqa: E402
from pyfpa.excel.model_workbook import model_to_excel  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
DRILL = REPO / "examples" / "keldagrim-forge" / "drill.py"
CONFIG = REPO / "examples" / "keldagrim-forge" / "config.yaml"
SCENARIOS = REPO / "examples" / "keldagrim-forge" / "scenarios.yaml"


def _run_drill(tmp_path, mutate=None, scenario="base"):
    cfg = load_config(CONFIG)
    path = tmp_path / "submission.xlsx"
    model_to_excel(cfg, path)
    if mutate is not None:
        from openpyxl import load_workbook
        wb = load_workbook(path)
        mutate(wb)
        wb.save(path)
    return subprocess.run(
        [sys.executable, str(DRILL), str(path), "--scenario", scenario],
        capture_output=True, text=True, cwd=REPO, timeout=300,
    )


def _payload(result):
    # the CLI contract: stdout is one JSON document
    assert result.stdout.strip(), result.stderr
    return json.loads(result.stdout)


def _labels(model):
    return {model.cell(row=r, column=1).value: r for r in range(2, model.max_row + 1)}


def test_drill_passes_on_a_faithful_workbook(tmp_path):
    result = _run_drill(tmp_path)
    payload = _payload(result)
    assert result.returncode == 0, result.stderr
    assert payload["structure"]["passed"] is True
    assert payload["numbers"]["passed"] is True
    assert payload["verdict"] == "PASS"


def test_drill_fails_when_a_driver_is_typed_over_a_formula(tmp_path):
    def mutate(wb):
        model = wb["Model"]
        model.cell(row=_labels(model)["revenue"], column=3, value=420000.0)

    result = _run_drill(tmp_path, mutate=mutate)
    payload = _payload(result)
    assert result.returncode == 1
    assert payload["structure"]["passed"] is False
    assert payload["verdict"] == "FAIL"


def test_drill_fails_when_the_numbers_diverge(tmp_path):
    # a wrong formula holds the structure together (the self-checks still tie)
    # but the engine comparison catches the divergence
    def mutate(wb):
        model = wb["Model"]
        row = _labels(model)["revenue"]
        model.cell(row=row, column=3, value=f"=B{row}*1.5")

    result = _run_drill(tmp_path, mutate=mutate)
    payload = _payload(result)
    assert result.returncode == 1
    assert payload["structure"]["passed"] is True
    assert payload["numbers"]["passed"] is False


def test_drill_fails_when_a_check_row_is_broken(tmp_path):
    def mutate(wb):
        model = wb["Model"]
        checks = [r for label, r in _labels(model).items()
                  if isinstance(label, str) and label.startswith("check_")]
        model.cell(row=checks[0], column=3, value="=123")

    result = _run_drill(tmp_path, mutate=mutate)
    payload = _payload(result)
    assert result.returncode == 1
    assert payload["numbers"]["passed"] is False
    assert any("non-zero" in f for f in payload["numbers"]["failures"])


def test_receipt_delay_scenario_scores_the_delayed_workbook(tmp_path):
    shift = yaml.safe_load(SCENARIOS.read_text(encoding="utf-8"))["receipt-delay"]
    amount = float(shift["amount"])

    def delay(wb):
        model = wb["Model"]
        row = _labels(model)["wc_cash_impact"]
        columns = {
            model.cell(row=1, column=c).value: c
            for c in range(2, model.max_column + 1)
        }
        for label, sign in ((shift["month"], "-"), (shift["to_month"], "+")):
            column = columns[label]
            original = model.cell(row=row, column=column).value
            model.cell(row=row, column=column, value=f"={original[1:]}{sign}({amount})")

    # the undelayed build must fail the delayed scenario ...
    result = _run_drill(tmp_path, scenario="receipt-delay")
    payload = _payload(result)
    assert result.returncode == 1
    assert payload["numbers"]["passed"] is False

    # ... and the delayed build must pass it while failing the base scenario
    result = _run_drill(tmp_path, mutate=delay, scenario="receipt-delay")
    payload = _payload(result)
    assert result.returncode == 0, result.stderr
    assert payload["verdict"] == "PASS"

    result = _run_drill(tmp_path, mutate=delay, scenario="base")
    payload = _payload(result)
    assert result.returncode == 1
    assert payload["numbers"]["passed"] is False
