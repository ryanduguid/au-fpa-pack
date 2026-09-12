"""Regression guards for the Harbour Light synthetic Australian example.

Seams: Xero fixture mapping, quarterly BAS cash dates, and live-formula
workbook verification against the engine.
"""
import json
import subprocess
import sys
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "examples" / "harbour-light"
sys.path.insert(0, str(EXAMPLE))

pytest.importorskip("formulas")

from models.generated.harbour_excel import export_workbook


def test_channels_refuse_a_missing_sales_account(monkeypatch):
    from types import SimpleNamespace

    import harbour_model as hm

    report = hm.profit_and_loss()
    accounts = report.by_account()
    accounts.pop("Sales - GST Free")
    monkeypatch.setattr(hm, "profit_and_loss", lambda: SimpleNamespace(by_account=lambda: accounts, by_tracking=report.by_tracking))
    with pytest.raises(ValueError, match="required sales account"):
        hm.channels_from_xero()


def test_xero_mapping_annualises_tracking_channels():
    import harbour_model as hm

    channels = {channel.name: channel for channel in hm.channels_from_xero()}
    assert channels["North"].annual_revenue == pytest.approx(840_000.0)
    assert channels["South"].annual_revenue == pytest.approx(492_000.0)
    assert channels["North"].cogs_pct == pytest.approx(31_000.0 / 70_000.0)
    assert channels["South"].cogs_pct == pytest.approx(19_000.0 / 41_000.0)
    assert hm.taxable_sales_pct() == pytest.approx(99_000.0 / 111_000.0)


def test_payroll_gross_wages_tie_to_xero_wages_line():
    import harbour_model as hm

    payroll = hm.payroll_frame()
    assert payroll["gross_wages"].iloc[0] == pytest.approx(41_600.0)
    assert payroll["super_guarantee"].iloc[0] == pytest.approx(4_992.0)
    # Synthetic Xero has $1,300 payroll tax; VIC wages sit under the $1m
    # threshold, so the kernel is the statutory source and pays nil.
    assert payroll["payroll_tax"].sum() == pytest.approx(0.0)


def test_quarterly_bas_dates_match_ato_fy2027_cycle():
    import harbour_model as hm

    schedule = hm.bas_settlement()
    assert schedule["due_date"].tolist() == [
        date(2026, 7, 28),  # opening GST balances, synthetic settlement assumption
        date(2026, 10, 28),
        date(2027, 2, 28),
        date(2027, 4, 28),
        date(2027, 7, 28),
    ]
    receipts, disbursements = hm.gst_cash13_flows()
    assert receipts == []
    assert len(disbursements) == 1
    assert disbursements[0].start_week == 4
    assert disbursements[0].name == "BAS 2027Q1"


def test_fy_summary_labels_harbour_months_as_fy2027():
    import harbour_model as hm

    from pyfpa.au import fy_summary

    monthly = hm.monthly_forecast()
    quarters = fy_summary(monthly[["revenue"]], by="quarter")
    assert list(quarters.index) == [
        "Q1 FY2027",
        "Q2 FY2027",
        "Q3 FY2027",
        "Q4 FY2027",
    ]
    assert quarters["revenue"].sum() == pytest.approx(1_332_000.0)


def test_verified_excel_matches_engine(tmp_path):
    import harbour_model as hm

    from pyfpa.excel.verify import verify_workbook

    path = tmp_path / "harbour-model.xlsx"
    export_workbook(path)
    monthly = hm.monthly_forecast()
    report = verify_workbook(path, monthly)
    assert report.passed, report.failures
    assert report.lines_checked == len(monthly.columns)


def test_monthly_cash_preserves_leave_expense_and_reconciles_gst():
    import harbour_model as hm

    monthly = hm.monthly_forecast()
    # Monthly gross profit 61,000 less overhead 16,900 and payroll cost 51,331.20.
    assert monthly["net_income"].sum() == pytest.approx(-86_774.40)
    # Leave 46,886.40 is non-cash. GST accrued 38,520 less 32,740 settled.
    assert monthly["ending_cash"].iloc[-1] == pytest.approx(51_312.0)
    assert monthly["leave_provisions"].sum() == pytest.approx(46_886.40)
    assert monthly["gst_accrued"].sum() == pytest.approx(38_520.0)
    assert monthly["gst_settled"].sum() == pytest.approx(32_740.0)
    assert monthly["gst_payable"].iloc[-1] == pytest.approx(9_630.0)
    assert monthly["ending_cash"].iloc[2] == pytest.approx(81_228.0)


def test_weekly_bank_movements_include_gst_and_use_reconciled_opening_cash():
    import harbour_model as hm

    from pyfpa.cash13.forecast import cash13_forecast

    cfg = hm.cash13_config()
    assert cfg.opening_cash == pytest.approx(81_228.0)
    receipts = {flow.name: flow.amount for flow in cfg.receipts}
    payments = {flow.name: flow.amount for flow in cfg.disbursements}
    assert receipts["Collections"] == pytest.approx(27_900.0)
    assert payments["Supplier payments"] == pytest.approx(165_000.0 / 13)
    assert payments["Operating overhead"] == pytest.approx(4_290.0)
    assert payments["BAS 2027Q1"] == pytest.approx(9_630.0)
    assert cash13_forecast(cfg)["ending_cash"].iloc[-1] == pytest.approx(60_312.0)


def test_run_harbour_writes_artifacts(tmp_path):
    import run_harbour as runner

    result = runner.run_harbour(tmp_path)
    assert (tmp_path / "briefing.md").exists()
    assert (tmp_path / "model.xlsx").exists()
    assert result["revenue_total"] == 1_332_000
    text = (tmp_path / "briefing.md").read_text()
    assert "# Harbour Light Pty Ltd" in text
    assert "13-Week" in text


def test_workbook_cash_recalculates_after_opening_gst_edit(tmp_path):
    import harbour_model as hm
    from openpyxl import load_workbook

    from pyfpa.excel.verify import verify_workbook

    path = tmp_path / "edited.xlsx"
    export_workbook(path)
    wb = load_workbook(path)
    sheet, address = next(wb.defined_names["opening_gst"].destinations)
    wb[sheet][address] = 4_850.0
    wb.save(path)
    expected = hm.monthly_forecast()
    expected.loc[expected.index[0], "gst_settled"] += 1_000.0
    for line in ("gst_cash_impact", "operating_cash_flow", "free_cash_flow", "change_in_cash"):
        expected.loc[expected.index[0], line] -= 1_000.0
    expected["ending_cash"] -= 1_000.0
    report = verify_workbook(path, expected)
    assert report.passed, report.failures


def test_export_verification_failure_preserves_existing_delivery(tmp_path, monkeypatch):
    import harbour_model as hm

    path = tmp_path / "model.xlsx"
    path.write_bytes(b"previous verified delivery")
    wrong_expected = hm.monthly_forecast()
    wrong_expected["ending_cash"] += 1_000.0
    monkeypatch.setattr(hm, "monthly_forecast", lambda: wrong_expected)
    with pytest.raises(ValueError, match="verification failed"):
        export_workbook(path)
    assert path.read_bytes() == b"previous verified delivery"
    assert list(tmp_path.iterdir()) == [path]


def test_runner_rejects_corrupted_workbook_before_writing_briefing(tmp_path, monkeypatch):
    import run_harbour as runner
    from models.generated import harbour_excel
    from openpyxl import load_workbook

    export = harbour_excel._build_workbook

    def corrupt(path):
        export(path)
        wb = load_workbook(path)
        ws = wb["Model"]
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "ending_cash")
        ws.cell(row, 2, "=0")
        wb.save(path)

    monkeypatch.setattr(harbour_excel, "_build_workbook", corrupt)
    with pytest.raises(ValueError, match="verification failed"):
        runner.run_harbour(tmp_path)
    assert not (tmp_path / "briefing.md").exists()
    assert not (tmp_path / "model.xlsx").exists()


def test_harbour_pipeline_is_registered_for_agent_discovery():
    from pyfpa.memory.entrypoints import load_entrypoint_registry

    registry = load_entrypoint_registry(EXAMPLE / ".fpa" / "models" / "entrypoints.yaml")
    entrypoint = next(item for item in registry.entrypoints if item.name == "harbour-light-pipeline")
    assert entrypoint.kind == "forecast"
    assert entrypoint.command == ["python3", "run_harbour.py"]


def test_harbour_pl_mapping_covers_every_source_row():
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "pyfpa.cli",
            "reconcile-source",
            str(EXAMPLE),
            "--source-id",
            "xero-pl",
            "--account-column",
            "Account",
            "--amount-column",
            "Amount",
        ],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    assert result.returncode == 1, result.stdout
    evidence = json.loads(result.stdout)["data"]
    assert evidence["unmapped"] == []
    assert evidence["duplicates"] == []
    assert evidence["expected_provided"] is False


def test_harbour_workspace_passes_agent_toolbelt_diagnostics():
    result = subprocess.run(
        [sys.executable, "-m", "pyfpa.cli", "doctor", str(EXAMPLE)],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    assert result.returncode == 0, result.stdout


def test_registered_report_exports_verified_workbook(tmp_path):
    from pyfpa.memory.entrypoints import load_entrypoint_registry

    registry = load_entrypoint_registry(EXAMPLE / ".fpa" / "models" / "entrypoints.yaml")
    entrypoint = next(item for item in registry.entrypoints if item.kind == "report")
    path = tmp_path / "model.xlsx"
    result = subprocess.run(
        [sys.executable, *entrypoint.command[1:], "--output", str(path)],
        cwd=EXAMPLE, text=True, capture_output=True, check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert path.exists()
