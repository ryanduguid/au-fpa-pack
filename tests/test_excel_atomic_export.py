import os
import stat
import zipfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from pyfpa.excel.model_workbook import model_to_excel
from pyfpa.io.reporting import forecast_to_excel

REPO = Path(__file__).resolve().parents[1]


def export_model(path):
    from pyfpa.config.loader import load_config

    model_to_excel(load_config(REPO / "examples/ridgeline/config.yaml"), path)


def export_forecast(path):
    forecast_to_excel(pd.DataFrame({"revenue": [12.5]}), path)


@pytest.mark.parametrize("export", [export_model, export_forecast])
@pytest.mark.parametrize("existing", [False, True])
def test_failed_save_preserves_destination_and_removes_staging(tmp_path, monkeypatch, export, existing):
    destination = tmp_path / "report.xlsx"
    if existing:
        destination.write_bytes(b"previous report")
    original = set(tmp_path.iterdir())

    def fail_save(self, path, **kwargs):
        path.write(b"incomplete archive")
        raise OSError("disk full")

    if export is export_forecast:
        monkeypatch.setattr(pd.DataFrame, "to_excel", fail_save)
    else:
        monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError, match="disk full"):
        export(destination)
    assert set(tmp_path.iterdir()) == original
    if existing:
        assert destination.read_bytes() == b"previous report"


@pytest.mark.parametrize("export", [export_model, export_forecast])
def test_success_replaces_existing_workbook_and_cleans_staging(tmp_path, export):
    destination = tmp_path / "report.xlsx"
    destination.write_bytes(b"previous report")
    export(destination)
    workbook = load_workbook(destination)
    try:
        assert workbook.sheetnames in (["Assumptions", "Model"], ["Forecast"])
    finally:
        workbook.close()
    assert list(tmp_path.iterdir()) == [destination]


@pytest.mark.skipif(os.name == "nt", reason="POSIX mode bits are checked on Linux")
def test_replacement_retains_restricted_permissions(tmp_path):
    destination = tmp_path / "report.xlsx"
    destination.write_bytes(b"previous report")
    destination.chmod(0o600)
    export_forecast(destination)
    assert stat.S_IMODE(destination.stat().st_mode) == 0o600


@pytest.mark.skipif(os.name == "nt", reason="Windows symlinks can require elevated privileges")
def test_replacement_keeps_existing_symlink(tmp_path):
    target = tmp_path / "actual.xlsx"
    target.write_bytes(b"previous report")
    destination = tmp_path / "report.xlsx"
    destination.symlink_to(target)
    export_forecast(destination)
    assert destination.is_symlink()
    workbook = load_workbook(target)
    try:
        assert workbook.sheetnames == ["Forecast"]
    finally:
        workbook.close()
    assert set(tmp_path.iterdir()) == {target, destination}


@pytest.mark.parametrize("export", [export_model, export_forecast])
def test_failed_replacement_preserves_previous_report(tmp_path, monkeypatch, export):
    destination = tmp_path / "report.xlsx"
    destination.write_bytes(b"previous report")

    def fail_replace(source, target):
        assert Path(source).is_file()
        raise PermissionError("workbook is open")

    monkeypatch.setattr("pyfpa.io.atomic.os.replace", fail_replace)
    with pytest.raises(PermissionError, match="workbook is open"):
        export(destination)
    assert destination.read_bytes() == b"previous report"
    assert list(tmp_path.iterdir()) == [destination]


def test_forecast_refuses_non_xlsx_destination(tmp_path):
    destination = tmp_path / "report.ods"
    with pytest.raises(ValueError, match=r"\.xlsx only"):
        export_forecast(destination)
    assert list(tmp_path.iterdir()) == []


def test_another_hard_link_retains_previous_report(tmp_path):
    destination = tmp_path / "report.xlsx"
    previous = tmp_path / "previous.xlsx"
    destination.write_bytes(b"previous report")
    os.link(destination, previous)
    export_forecast(destination)
    assert previous.read_bytes() == b"previous report"
    workbook = load_workbook(destination)
    try:
        assert workbook.sheetnames == ["Forecast"]
    finally:
        workbook.close()


def test_staged_verification_refuses_incomplete_archive(tmp_path, monkeypatch):
    destination = tmp_path / "report.xlsx"
    destination.write_bytes(b"previous report")

    def incomplete_save(self, stream):
        stream.write(b"incomplete archive")

    monkeypatch.setattr(Workbook, "save", incomplete_save)
    with pytest.raises(zipfile.BadZipFile):
        export_model(destination)
    assert destination.read_bytes() == b"previous report"
    assert list(tmp_path.iterdir()) == [destination]
