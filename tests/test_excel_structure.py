"""Structure checks for model workbooks: layout contract, no recalculation.

verify_workbook proves the numbers. These checks prove the build standard:
inputs styled and confined to Assumptions, formula-only Model rows, a
restricted formula vocabulary, aligned period headers and a checks row.
"""

from pyfpa.config.schemas import EntityConfig
from pyfpa.excel.model_workbook import model_to_excel
from pyfpa.excel.structure import verify_structure
from pyfpa.excel.toolkit import INPUT_FONT_COLOR


def _cfg():
    return EntityConfig.model_validate({
        "name": "T", "start_month": "2026-01", "horizon_months": 6, "tax_rate": 0.21,
        "channels": [
            {"name": "A", "annual_revenue": 1_200_000.0, "growth_rate": 0.0,
             "seasonality": [1.0] * 12, "cogs_pct": 0.5},
        ],
        "opex": [{"name": "salaries", "kind": "fixed", "monthly_amount": 30_000.0}],
        "debt": [],
        "working_capital": {"dso_days": 30.0, "dpo_days": 30.0, "dio_days": 30.0},
        "opening_balances": {"cash": 10_000.0},
        "da_monthly": 1_000.0, "capex_monthly": 500.0,
    })


def _build(tmp_path):
    path = tmp_path / "model.xlsx"
    model_to_excel(_cfg(), path)
    return path


def test_faithful_export_passes_structure(tmp_path):
    report = verify_structure(_build(tmp_path))
    assert report.passed, report.failures
    assert report.checks_run > 0


def test_missing_model_sheet_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    del wb["Model"]
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("Model" in f and "sheet" in f for f in report.failures)


def test_missing_assumptions_sheet_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    del wb["Assumptions"]
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("Assumptions" in f and "sheet" in f for f in report.failures)


def test_value_typed_over_a_formula_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    model = wb["Model"]
    labels = {model.cell(row=r, column=1).value: r for r in range(2, model.max_row + 1)}
    model.cell(row=labels["revenue"], column=3, value=123456.0)
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("revenue" in f and "formula" in f for f in report.failures)


def test_forbidden_function_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    model = wb["Model"]
    labels = {model.cell(row=r, column=1).value: r for r in range(2, model.max_row + 1)}
    model.cell(row=labels["revenue"], column=3, value="=LOG10(100)")
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("LOG10" in f for f in report.failures)


def test_empty_or_duplicate_period_header_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    model = wb["Model"]
    model.cell(row=1, column=3).value = None
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("header" in f.lower() for f in report.failures)

    path = _build(tmp_path)
    wb = load_workbook(path)
    model = wb["Model"]
    first = model.cell(row=1, column=2).value
    model.cell(row=1, column=4, value=first)
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("duplicate" in f.lower() for f in report.failures)


def test_unlabelled_model_row_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    model = wb["Model"]
    labels = {model.cell(row=r, column=1).value: r for r in range(2, model.max_row + 1)}
    model.cell(row=labels["revenue"], column=1).value = None
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("label" in f.lower() for f in report.failures)


def test_defined_name_on_the_model_sheet_fails(tmp_path):
    from openpyxl import load_workbook
    from openpyxl.workbook.defined_name import DefinedName
    path = _build(tmp_path)
    wb = load_workbook(path)
    wb.defined_names["sneaky"] = DefinedName(name="sneaky", attr_text="'Model'!$B$2")
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("sneaky" in f and "Assumptions" in f for f in report.failures)


def test_unstyled_input_cell_fails(tmp_path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    path = _build(tmp_path)
    wb = load_workbook(path)
    ws = wb["Assumptions"]
    # strip the input colour from the first named driver cell
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)
        if cell.value is not None and not isinstance(cell.value, str):
            cell.font = Font(color="FF000000")
            break
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("input" in f.lower() and "colour" in f.lower() for f in report.failures)


def test_input_colour_on_a_model_cell_fails(tmp_path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    path = _build(tmp_path)
    wb = load_workbook(path)
    model = wb["Model"]
    labels = {model.cell(row=r, column=1).value: r for r in range(2, model.max_row + 1)}
    model.cell(row=labels["revenue"], column=3).font = Font(color=INPUT_FONT_COLOR)
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("Model" in f and "input" in f.lower() for f in report.failures)


def test_missing_checks_row_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    model = wb["Model"]
    for r in range(2, model.max_row + 1):
        label = model.cell(row=r, column=1).value
        if isinstance(label, str) and label.startswith("check_"):
            model.cell(row=r, column=1, value=label.replace("check_", "tie_"))
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("check_" in f for f in report.failures)


def test_formula_on_an_input_cell_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    ws = wb["Assumptions"]
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)
        if cell.value is not None and not isinstance(cell.value, str):
            cell.value = "=1+1"
            break
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("input" in f.lower() and "formula" in f.lower() for f in report.failures)


def test_non_numeric_input_cell_fails(tmp_path):
    from openpyxl import load_workbook
    path = _build(tmp_path)
    wb = load_workbook(path)
    ws = wb["Assumptions"]
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)
        if cell.value is not None and not isinstance(cell.value, str):
            cell.value = None
            break
    wb.save(path)
    report = verify_structure(path)
    assert not report.passed
    assert any("is not a number" in f for f in report.failures)
