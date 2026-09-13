"""Synthetic NSW service business: dated cash, monthly profit and verified Excel."""
from __future__ import annotations

import argparse
import json
from datetime import timedelta
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from pyfpa.au.payroll import PayrollAssumptions, Role, payroll_forecast
from pyfpa.cash13.forecast import cash13_forecast
from pyfpa.cash13.schemas import Cash13Config, WeeklyFlow
from pyfpa.excel.toolkit import add_named_cell
from pyfpa.excel.verify import verify_workbook
from pyfpa.io.xero_au import read_xero_report

ROOT = Path(__file__).resolve().parents[2]
MONTHS = pd.period_range("2026-10", "2026-12", freq="M")
PL_ACCOUNTS = {
    "Sales - Varrock maintenance", "Sales - Falador repairs", "Materials",
    "Operating overhead", "Wages and Salaries", "Superannuation",
    "Workers Compensation", "Leave provision expense", "Depreciation", "Interest expense",
}
BS_ACCOUNTS = {
    "Business Bank Account", "Accounts Receivable", "Equipment net book value",
    "Accounts Payable", "GST", "PAYG Withholdings Payable", "Business Loan", "Equity",
}
MONEY = '#,##0.00;[Red](#,##0.00);"-"'


def forecast(data: Path = ROOT / "data", delay_days: int = 0) -> dict:
    if not isinstance(delay_days, int) or not 0 <= delay_days <= 120:
        raise ValueError("Receipt delay must be an integer from 0 to 120 days")
    a = json.loads((data / "assumptions.json").read_text(encoding="utf-8"))
    pl = read_xero_report(data / "xero_pl.csv").by_account()
    bs = read_xero_report(data / "xero_bs.csv").by_account()
    if set(pl) != PL_ACCOUNTS or set(bs) != BS_ACCOUNTS:
        raise ValueError("Missing or unmapped source account")
    if abs(sum(bs.values())) > 0.01:
        raise ValueError("Opening balance sheet does not balance")
    roles = [Role(**row) for row in pd.read_csv(data / "payroll.csv").to_dict("records")]
    payroll = payroll_forecast(roles, MONTHS, PayrollAssumptions(
        workers_comp_rate=a["workers_comp_rate"], annual_leave_weeks=a["annual_leave_weeks"],
        lsl_accrual_pct=a["lsl_accrual_pct"],
    )).round(2)
    for source, column in [("Wages and Salaries", "gross_wages"), ("Superannuation", "super_guarantee"),
                           ("Workers Compensation", "workers_comp"), ("Leave provision expense", "leave_provisions")]:
        if not (abs(payroll[column] + pl[source]) < 0.01).all():
            raise ValueError(f"Payroll does not reconcile to source account {source}")
    if payroll["payroll_tax"].any():
        raise ValueError("This example needs a payroll-tax payment schedule above the NSW threshold")
    invoices = pd.read_csv(data / "invoices.csv", parse_dates=["receipt_date"])
    payments = pd.read_csv(data / "payments.csv", parse_dates=["date"])
    if invoices["receipt_date"].isna().any() or payments["date"].isna().any():
        raise ValueError("Missing cash date")
    if set(invoices["service_month"]) != {"2026-09", *MONTHS.astype(str)} or not payments["date"].dt.to_period("M").isin(MONTHS).all():
        raise ValueError("Unexpected source period")
    if (invoices["receipt_date"] < pd.Timestamp(a["forecast_start"])).any():
        raise ValueError("Opening and forecast invoices cannot be collected before the forecast")
    numeric = pd.concat([invoices["net"], invoices["gst"], payments["amount"]])
    if numeric.isna().any() or not numeric.map(lambda x: 0 <= x < float("inf")).all():
        raise ValueError("Cash source amounts must be finite and non-negative")
    if invoices["id"].duplicated().any() or not (invoices["id"] == a["delay_invoice"]).sum() == 1:
        raise ValueError("Invoice identifiers must be unique and include the delayed invoice")
    if not (abs(invoices["gst"] - invoices["net"] * a["gst_rate"]) < 0.01).all():
        raise ValueError("Invoice GST does not reconcile")
    for month in ("2026-09", *MONTHS.astype(str)):
        rows = invoices[invoices["service_month"] == month]
        actual = rows.groupby("service_line")["net"].sum().to_dict()
        if actual != {"Varrock": pl["Sales - Varrock maintenance"], "Falador": pl["Sales - Falador repairs"]}:
            raise ValueError(f"Invoice revenue does not reconcile for {month}")
    opening = invoices[invoices["service_month"] == "2026-09"]
    if abs(opening[["net", "gst"]].sum().sum() - bs["Accounts Receivable"]) > 0.01:
        raise ValueError("Opening receivables do not reconcile")
    # ponytail: this fixed 3-month case refuses changed payment totals; extend the
    # liability roll-forwards before using it for a different supplier or payroll cycle.
    expected_payments = {
        "Materials": -pl["Materials"] * (1 + a["gst_rate"]),
        "Overhead": -pl["Operating overhead"] * (1 + a["gst_rate"]),
        "Net wages": -pl["Wages and Salaries"] - a["monthly_payg_withheld"],
        "Super": -pl["Superannuation"], "Workers compensation": -pl["Workers Compensation"],
        "PAYG withholding": a["monthly_payg_withheld"], "Interest": -pl["Interest expense"],
        "Loan principal": a["monthly_loan_principal"],
    }
    if abs(expected_payments["Materials"] + bs["Accounts Payable"]) > 0.01 or abs(a["monthly_payg_withheld"] + bs["PAYG Withholdings Payable"]) > 0.01:
        raise ValueError("Opening supplier or withholding liability does not reconcile")
    for month in MONTHS:
        expected = dict(expected_payments)
        if month == MONTHS[0]:
            expected.update({"GST settlement": -bs["GST"], "Income tax instalment": a["quarterly_tax_instalment"]})
        actual = payments[payments["date"].dt.to_period("M") == month].groupby("category")["amount"].sum().to_dict()
        if set(actual) != set(expected) or any(abs(actual[k] - v) > 0.01 for k, v in expected.items()):
            raise ValueError(f"Payment schedule does not reconcile for {month}")
    invoices["base_date"] = invoices["receipt_date"]
    invoices.loc[invoices["id"] == a["delay_invoice"], "receipt_date"] += timedelta(days=delay_days)
    receipts = pd.DataFrame({"id": invoices["id"], "base_date": invoices["base_date"], "date": invoices["receipt_date"],
                             "category": "Customer receipt", "receipts": invoices["net"] + invoices["gst"], "payments": 0.0})
    outgoing = pd.DataFrame({"id": [f"PAY-{i + 1:02}" for i in range(len(payments))], "base_date": payments["date"],
                            "date": payments["date"], "category": payments["category"], "receipts": 0.0, "payments": payments["amount"]})
    ledger = pd.concat([receipts, outgoing], ignore_index=True).sort_values("base_date").reset_index(drop=True)
    start, end = pd.Timestamp(a["weekly_start"]), pd.Timestamp(a["forecast_end"])
    if ((ledger["date"] < start) & (ledger["date"] >= pd.Timestamp(a["forecast_start"]))).any():
        raise ValueError("Cash on 1 October falls outside the 13-week window")
    flows = ledger[ledger["date"].between(start, end)]
    weekly = cash13_forecast(Cash13Config(
        opening_cash=bs["Business Bank Account"],
        receipts=[WeeklyFlow(name=r.id, amount=r.receipts, start_week=(r.date - start).days // 7 + 1) for r in flows.itertuples() if r.receipts],
        disbursements=[WeeklyFlow(name=r.id, amount=r.payments, start_week=(r.date - start).days // 7 + 1) for r in flows.itertuples() if r.payments],
    ))
    daily = flows.groupby("date")[["receipts", "payments"]].sum().reindex(pd.date_range(start, end), fill_value=0)
    daily["ending_cash"] = bs["Business Bank Account"] + (daily["receipts"] - daily["payments"]).cumsum()
    monthly = pd.DataFrame(index=MONTHS)
    for label, value in {
        "Revenue": pl["Sales - Varrock maintenance"] + pl["Sales - Falador repairs"],
        "Materials": -pl["Materials"], "Overhead": -pl["Operating overhead"],
        "Gross wages": -pl["Wages and Salaries"], "Super": -pl["Superannuation"],
        "Workers compensation": -pl["Workers Compensation"], "Leave provision": -pl["Leave provision expense"],
        "Depreciation": -pl["Depreciation"], "Interest": -pl["Interest expense"],
    }.items():
        monthly[label] = value
    monthly["Profit before tax"] = monthly["Revenue"] - monthly.drop(columns="Revenue").sum(axis=1)
    groups = flows.groupby(flows["date"].dt.to_period("M"))[["receipts", "payments"]].sum().reindex(MONTHS, fill_value=0)
    monthly["Customer receipts"] = groups["receipts"]
    monthly["Cash payments"] = groups["payments"]
    monthly["Closing cash"] = bs["Business Bank Account"] + (groups["receipts"] - groups["payments"]).cumsum()
    monthly["Receivables"] = bs["Accounts Receivable"] + (monthly["Revenue"] * (1 + a["gst_rate"]) - groups["receipts"]).cumsum()
    monthly["Supplier payables"] = -bs["Accounts Payable"]
    settlements = payments[payments["category"] == "GST settlement"].groupby(payments["date"].dt.to_period("M"))["amount"].sum().reindex(MONTHS, fill_value=0)
    monthly["GST payable"] = -bs["GST"] + ((monthly["Revenue"] - monthly["Materials"] - monthly["Overhead"]) * a["gst_rate"] - settlements).cumsum()
    monthly["PAYG withholding payable"] = a["monthly_payg_withheld"]
    monthly["Loan principal paid"] = a["monthly_loan_principal"]
    monthly["Income tax instalment"] = [a["quarterly_tax_instalment"], 0, 0]
    ar_change = monthly["Receivables"].diff().fillna(monthly["Receivables"].iloc[0] - bs["Accounts Receivable"])
    gst_change = monthly["GST payable"].diff().fillna(monthly["GST payable"].iloc[0] + bs["GST"])
    monthly["Cash from profit bridge"] = monthly["Profit before tax"] + monthly["Depreciation"] + monthly["Leave provision"] - ar_change + gst_change - monthly["Loan principal paid"] - monthly["Income tax instalment"]
    monthly["Cash reconciliation difference"] = monthly["Cash from profit bridge"] - (groups["receipts"] - groups["payments"])
    if monthly["Cash reconciliation difference"].abs().max() >= 0.01:
        raise ValueError("Profit-to-cash reconciliation failed")
    return {"monthly": monthly.round(2), "weekly": weekly, "daily": daily, "ledger": ledger,
            "pl": pl, "bs": bs, "assumptions": a, "delay_days": delay_days, "data": data}


def export_workbook(path: Path, result: dict) -> None:
    """Use the repository toolkit; replace output only after formula verification."""
    wb = Workbook()
    wb.active.title = "Model"
    for name in ("Cash 13 weeks", "Assumptions", "Cash dates", "Sources"):
        wb.create_sheet(name)
    control = wb["Assumptions"]
    control.append(["Lumbridge Services: synthetic Newcastle business", "Value"])
    control.append(["Delay OPEN-V receipt (calendar days)", result["delay_days"]])
    add_named_cell(wb, control, name="receipt_delay", row=2, col=2, value=result["delay_days"])
    control.append(["Cash buffer (AUD)", result["assumptions"]["minimum_cash_buffer"]])
    control.append(["Change B2 to 45 to test a late Grand Exchange collection."])
    control.append(["Only B2 is an interactive scenario control. Rebuild after changing source data."])
    control.append(["All amounts are AUD. No game currency or live Xero connection."])
    control.append(["Dates move by calendar days; no holiday or receipt-compliance guarantee."])
    validation = DataValidation(type="whole", operator="between", formula1=0, formula2=120)
    validation.errorTitle, validation.error = "Invalid delay", "Enter a whole number from 0 to 120."
    validation.showErrorMessage, validation.errorStyle = True, "stop"
    control.add_data_validation(validation)
    validation.add(control["B2"])
    control["B2"].fill = PatternFill("solid", fgColor="FFF0C2")
    control["B2"].font = Font(color="0000FF", bold=True)
    source = wb["Sources"]
    source.append(["Synthetic source record", "Value", "Source / period"])
    refs = {}
    for kind in ("pl", "bs"):
        for label, value in result[kind].items():
            source.append([label, value, f"data/xero_{kind}.csv | September 2026"])
            refs[label] = f"Sources!$B${source.max_row}"
    source.append(["GST rate", result["assumptions"]["gst_rate"], "data/assumptions.json; see README source notes"])
    gst_ref = f"Sources!$B${source.max_row}"
    source.cell(source.max_row, 2).number_format = "0.0%"
    source.append(["Loan principal per month", result["assumptions"]["monthly_loan_principal"], "data/assumptions.json"])
    principal_ref = f"Sources!$B${source.max_row}"
    source.append(["Income tax instalment in October", result["assumptions"]["quarterly_tax_instalment"], "Synthetic supplied instalment, not a tax calculation"])
    tax_ref = f"Sources!$B${source.max_row}"
    cash = wb["Cash dates"]
    cash.append(["Record", "Base date", "Scenario date", "Receipts AUD", "Payments AUD", "Category", "Closing cash on date"])
    for i, event in enumerate(result["ledger"].itertuples(), 2):
        date_formula = f'=B{i}+IF(A{i}="OPEN-V",receipt_delay,0)'
        cash.append([event.id, event.base_date.to_pydatetime(), date_formula, event.receipts, event.payments, event.category])
        cash.cell(i, 2).number_format = cash.cell(i, 3).number_format = "dd-mmm-yy"
    last = cash.max_row
    dates = f"'Cash dates'!$C$2:$C${last}"
    receipts = f"'Cash dates'!$D$2:$D${last}"
    payments = f"'Cash dates'!$E$2:$E${last}"
    categories = f"'Cash dates'!$F$2:$F${last}"
    opening_cash = refs["Business Bank Account"]
    for i in range(2, last + 1):
        cash.cell(i, 7, f'=IF(C{i}>DATE(2026,12,31),"",{opening_cash}+SUMIFS({receipts},{dates},"<="&C{i})-SUMIFS({payments},{dates},"<="&C{i}))')
    ws = wb["Model"]
    ws.append(["Lumbridge Services: monthly forecast (AUD)", *[m.start_time.to_pydatetime() for m in MONTHS]])
    rows = {name: i for i, name in enumerate(result["monthly"].columns, 2)}
    for name, row in rows.items():
        ws.cell(row, 1, name)
    input_lines = {"Materials": "Materials", "Overhead": "Operating overhead", "Gross wages": "Wages and Salaries",
                   "Super": "Superannuation", "Workers compensation": "Workers Compensation", "Leave provision": "Leave provision expense",
                   "Depreciation": "Depreciation", "Interest": "Interest expense"}
    for i, month in enumerate(MONTHS, 2):
        col, prev = get_column_letter(i), get_column_letter(i - 1)
        ref = lambda label, col=col: f"{col}{rows[label]}"
        prior = lambda label, initial, i=i, prev=prev: initial if i == 2 else f"{prev}{rows[label]}"
        first, end = month.start_time.to_pydatetime(), month.end_time.normalize().to_pydatetime()
        start_expr, end_expr = f"DATE({first.year},{first.month},1)", f"DATE({end.year},{end.month},{end.day})"
        cash_sum = lambda area, start_expr=start_expr, end_expr=end_expr: f'SUMIFS({area},{dates},">="&{start_expr},{dates},"<="&{end_expr})'
        formulas = {"Revenue": f"={refs['Sales - Varrock maintenance']}+{refs['Sales - Falador repairs']}"}
        formulas.update({name: f"=-{refs[source_name]}" for name, source_name in input_lines.items()})
        formulas.update({
            "Profit before tax": f"={ref('Revenue')}-SUM({ref('Materials')}:{ref('Interest')})",
            "Customer receipts": "=" + cash_sum(receipts), "Cash payments": "=" + cash_sum(payments),
            "Closing cash": f"={prior('Closing cash', opening_cash)}+{ref('Customer receipts')}-{ref('Cash payments')}",
            "Receivables": f"={prior('Receivables', refs['Accounts Receivable'])}+{ref('Revenue')}*(1+{gst_ref})-{ref('Customer receipts')}",
            "Supplier payables": f"=-{refs['Accounts Payable']}",
            "GST payable": f'={prior("GST payable", "-" + refs["GST"])}+({ref("Revenue")}-{ref("Materials")}-{ref("Overhead")})*{gst_ref}-SUMIFS({payments},{categories},"GST settlement",{dates},">="&{start_expr},{dates},"<="&{end_expr})',
            "PAYG withholding payable": f"=-{refs['PAYG Withholdings Payable']}",
            "Loan principal paid": f"={principal_ref}", "Income tax instalment": f"={tax_ref}" if i == 2 else "=0",
            "Cash from profit bridge": f"={ref('Profit before tax')}+{ref('Depreciation')}+{ref('Leave provision')}-({ref('Receivables')}-{prior('Receivables', refs['Accounts Receivable'])})+({ref('GST payable')}-{prior('GST payable', '-' + refs['GST'])})-{ref('Loan principal paid')}-{ref('Income tax instalment')}",
            "Cash reconciliation difference": f"={ref('Cash from profit bridge')}-{ref('Customer receipts')}+{ref('Cash payments')}",
        })
        for label, formula in formulas.items():
            ws.cell(rows[label], i, formula)
        ws.cell(1, i).number_format = "mmm-yy"
    ws.cell(ws.max_row + 2, 1, "Profit is before income tax. PAYG instalments reduce cash, not pre-tax profit.")
    ws.cell(ws.max_row + 1, 1, "GST, supplier and withholding balances remain payable after the forecast.")
    ws.cell(ws.max_row + 1, 1, "Select the receipt delay on Assumptions. No actual forecast accuracy has been measured.")
    ws.cell(ws.max_row + 1, 1, "Receipt delay (calendar days)")
    ws.cell(ws.max_row, 2, "=receipt_delay").number_format = "0"
    weekly = wb["Cash 13 weeks"]
    weekly.append(["Week start", "Week end", "Receipts AUD", "Payments AUD", "Net cash AUD", "Closing cash AUD", "Buffer AUD"])
    start = pd.Timestamp(result["assumptions"]["weekly_start"])
    for i in range(2, 15):
        day = start + pd.Timedelta(weeks=i - 2)
        weekly.append([day.to_pydatetime(), (day + pd.Timedelta(days=6)).to_pydatetime(),
                       f'=SUMIFS({receipts},{dates},">="&A{i},{dates},"<="&B{i})',
                       f'=SUMIFS({payments},{dates},">="&A{i},{dates},"<="&B{i})',
                       f"=C{i}-D{i}", f"={opening_cash}+E{i}" if i == 2 else f"=F{i-1}+E{i}", "=Assumptions!$B$3"])
        weekly.cell(i, 1).number_format = weekly.cell(i, 2).number_format = "dd-mmm-yy"
    weekly.cell(17, 1, "Minimum week-end cash")
    weekly.cell(17, 2, "=MIN(F2:F14)")
    weekly.cell(18, 1, "Minimum cash on a transaction date")
    # Future invoice dates retain receivables evidence but do not imply a January cash forecast.
    weekly.cell(18, 2, f"=MIN({opening_cash},'Cash dates'!G2:G{last})")
    weekly.cell(19, 1, "Week-end balances can hide a shortfall earlier in the week.")
    weekly.cell(20, 1, "Receipt delay (calendar days)")
    weekly.cell(20, 2, "=receipt_delay").number_format = "0"
    chart = LineChart()
    chart.title = "Bank balance and cash buffer (AUD)"
    chart.add_data(Reference(weekly, min_col=6, max_col=7, min_row=1, max_row=14), titles_from_data=True)
    chart.set_categories(Reference(weekly, min_col=2, min_row=2, max_row=14))
    chart.x_axis = DateAxis(axId=10, crossAx=100, majorTimeUnit="days", majorUnit=14)
    chart.x_axis.delete = chart.y_axis.delete = False
    chart.x_axis.tickLblPos, chart.y_axis.tickLblPos = "low", "nextTo"
    chart.x_axis.numFmt, chart.y_axis.numFmt = "dd-mmm", '#,##0;(#,##0)'
    chart.legend.position, chart.legend.overlay = "r", False
    for series in chart.series:
        series.smooth = False
    chart.width, chart.height = 23, 9
    weekly.add_chart(chart, "A22")
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B2"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation, sheet.page_setup.paperSize = "landscape", sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 1
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="253E35")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 34
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 21
            for cell in row:
                if cell.coordinate == "B2" and sheet.title == "Assumptions":
                    continue
                cell.font = Font(name="Calibri", size=11, color="20352D" if sheet.title in ("Model", "Cash 13 weeks") else "008000" if cell.data_type == "f" and "!" in str(cell.value) else "000000")
                if (isinstance(cell.value, (int, float)) or cell.data_type == "f") and cell.number_format == "General":
                    cell.number_format = MONEY
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 20
    ws.column_dimensions["A"].width = 43
    for row in (rows["Profit before tax"], rows["Closing cash"], rows["Cash reconciliation difference"]):
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor="E8EEE7")
            cell.font = Font(name="Calibri", size=11, bold=True, color="20352D")
    for cell in ws[rows["Cash reconciliation difference"]][1:]:
        cell.number_format = "0.00;(0.00);0.00"
    for row in range(rows["Cash reconciliation difference"] + 2, ws.max_row):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 30
        ws.cell(row, 1).alignment = Alignment(wrap_text=True)
    control.column_dimensions["A"].width = 91
    control["B2"].number_format = "0"
    source.column_dimensions["A"].width, source.column_dimensions["C"].width = 34, 66
    cash.column_dimensions["F"].width = 23
    cash.auto_filter.ref = f"A1:G{last}"
    for row in (17, 18, 19):
        weekly.row_dimensions[row].height = 34
        weekly.cell(row, 1).alignment = Alignment(wrap_text=True)
    weekly.column_dimensions["A"].width = 30
    weekly.merge_cells("A19:G19")
    weekly.print_area, ws.print_area = "A1:G39", f"A1:D{ws.max_row}"
    path = Path(path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=".lumbridge-", dir=path.parent) as directory:
        candidate = Path(directory) / path.name
        wb.save(candidate)
        report = verify_workbook(candidate, result["monthly"])
        if not report.passed or report.lines_checked != len(result["monthly"].columns):
            raise ValueError(f"Workbook verification failed: {report.failures}")
        candidate.replace(path)


def run(output: Path, delay_days: int = 0) -> dict:
    result = forecast(delay_days=delay_days)
    output.mkdir(parents=True, exist_ok=True)
    export_workbook(output / "lumbridge.xlsx", result)
    result["monthly"].to_csv(output / "monthly.csv", float_format="%.2f")
    result["weekly"].to_csv(output / "cash13.csv", float_format="%.2f")
    result["daily"].to_csv(output / "daily-cash.csv", float_format="%.2f", index_label="date")
    summary = {"delay_days": delay_days, "quarter_revenue": float(result["monthly"]["Revenue"].sum()),
               "quarter_profit_before_tax": round(float(result["monthly"]["Profit before tax"].sum()), 2),
               "december_cash": float(result["monthly"]["Closing cash"].iloc[-1]),
               "minimum_daily_cash": float(result["daily"]["ending_cash"].min()),
               "minimum_cash_date": str(result["daily"]["ending_cash"].idxmin().date()),
               "maximum_cash_reconciliation_difference": float(result["monthly"]["Cash reconciliation difference"].abs().max())}
    (output / "results.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    base, late = forecast(), forecast(delay_days=45)
    lines = [
        "# Lumbridge Services management briefing", "",
        ("Synthetic Newcastle maintenance business. October to December 2026, AUD. "
         "Varrock maintenance and Falador repairs are fictional service lines inspired by Old School RuneScape."), "",
        ("The base case earns $35,957.55 before income tax, but a late customer receipt can still leave the bank overdrawn. "
         "Treat the Grand Exchange collection date as an assumption until the customer confirms it."), "",
        "| Measure | On time | OPEN-V paid 45 days late |", "| --- | ---: | ---: |",
    ]
    for label, left, right in [
        ("Quarter revenue", base["monthly"]["Revenue"].sum(), late["monthly"]["Revenue"].sum()),
        ("Quarter profit before tax", base["monthly"]["Profit before tax"].sum(), late["monthly"]["Profit before tax"].sum()),
        ("October closing cash", base["monthly"]["Closing cash"].iloc[0], late["monthly"]["Closing cash"].iloc[0]),
        ("Minimum daily cash", base["daily"]["ending_cash"].min(), late["daily"]["ending_cash"].min()),
        ("December closing cash", base["monthly"]["Closing cash"].iloc[-1], late["monthly"]["Closing cash"].iloc[-1]),
    ]:
        lines.append(f"| {label} | ${left:,.2f} | ${right:,.2f} |")
    lines += ["", "## Monthly review", "",
              "| Month | Revenue | Profit before tax | Closing cash |",
              "| --- | ---: | ---: | ---: |"]
    for month, row in result["monthly"].iterrows():
        lines.append(f"| {month} | ${row['Revenue']:,.2f} | ${row['Profit before tax']:,.2f} | ${row['Closing cash']:,.2f} |")
    lines += ["", f"This monthly table uses a {delay_days}-day receipt delay. The comparison above always shows the two labelled cases.", "",
              ("October includes $12,600 of opening GST and a $3,000 income tax instalment. "
              "The late case first becomes negative on 28 October and reaches ($25,160) on 6 November. "
              "December cash catches up because the delayed receipt arrives on 28 November; that recovery does not fund the earlier shortfall."), "",
              ("At 31 December, $12,600 GST, $5,000 PAYG withholding and $13,200 supplier invoices remain payable. "
              "The loan balance is $17,000 and the quarter adds $6,762.45 of unused leave provisions. "
              "A positive closing bank balance does not settle those obligations."), "",
              "## Decisions for the owner", "",
              ("- Confirm the $44,000 OPEN-V receipt date. The 45-day delay needs $25,160 of additional cash at the trough, "
               "or $40,160 to preserve the assumed $15,000 buffer. No overdraft is assumed or arranged."),
              ("- Hold any dragon pickaxe upgrade, meaning discretionary equipment spending, until collection timing and funding are resolved. "
               "The forecast currently includes no capital expenditure."),
              "- Reforecast from actual receipts and payments each week. Investigate changes instead of treating the repeated September margin as earned growth.", "",
              "## Evidence and limits", "",
              ("Revenue is held at September's fabricated $60,000 monthly run rate; monthly pre-tax profit remains $11,985.85. "
              "There are no October to December actuals, measured forecast errors, client outcomes or claimed time savings. "
              "Income tax expense, contract enforceability, debt covenants and a full statutory balance sheet are outside this case. "
              "PAYG withholding, workers compensation, loan terms and tax instalments are supplied scenario assumptions. "
              "Read the example README before applying the workflow to another business."), ""]
    (output / "briefing.md").write_text("\n".join(lines), encoding="utf-8")
    return summary


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "output")
    parser.add_argument("--delay-days", type=int, default=0)
    args = parser.parse_args()
    print(json.dumps(run(args.output, args.delay_days), indent=2))
