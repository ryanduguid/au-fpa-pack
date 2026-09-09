"""Assemble the Harbour Light Pty Ltd synthetic Australian example.

Xero GST-exclusive fixtures (one month, tracking by region) are annualised
into the monthly engine. Statutory payroll and quarterly BAS cash live in
``pyfpa.au``; they are not re-typed from the Xero payroll-tax line.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pyfpa.au.calendar import fy_month_range
from pyfpa.au.gst import (
    BasCycle,
    GstAssumptions,
    bas_schedule,
    gst_weekly_flows,
    monthly_gst,
)
from pyfpa.au.payroll import PayrollAssumptions, Role, payroll_forecast
from pyfpa.cash13.schemas import Cash13Config, WeeklyFlow
from pyfpa.config.schemas import (
    Channel,
    EntityConfig,
    OpeningBalances,
    OpexLine,
    WorkingCapitalConfig,
)
from pyfpa.excel.model_workbook import model_to_excel
from pyfpa.excel.toolkit import MONEY_FORMAT, add_named_cell
from pyfpa.io.xero_au import read_xero_report
from pyfpa.models.cashflow import cashflow_from_config

DATA = Path(__file__).parent / "data"
CASH13_START = date(2026, 10, 1)
# Synthetic assumption: both opening GST accounts settle with the June BAS.
OPENING_GST_DUE = date(2026, 7, 28)
_DAYS_PER_YEAR = 360.0
_NON_PAYROLL_OPEX = (
    "Advertising",
    "Consulting & Accounting",
    "Office Expenses",
    "Rent",
    "Repairs & Maintenance",
)
# Gross wages tie to the Xero "Wages and Salaries" line ($41,600 / month).
ROLES = [
    Role(name="Workshop manager", annual_salary=130_000, jurisdiction="VIC"),
    Role(name="Electrician", annual_salary=95_000, jurisdiction="VIC"),
    Role(name="Sales", annual_salary=90_000, jurisdiction="VIC"),
    Role(name="Warehouse", annual_salary=85_000, jurisdiction="VIC"),
    Role(name="Bookkeeper", annual_salary=99_200, jurisdiction="VIC"),
]


def profit_and_loss():
    return read_xero_report(DATA / "xero_pl_tracking.csv")


def balance_sheet():
    return read_xero_report(DATA / "xero_bs.csv")


def taxable_sales_pct() -> float:
    accounts = profit_and_loss().by_account()
    domestic = accounts["Sales - Domestic"]
    gst_free = accounts["Sales - GST Free"]
    return domestic / (domestic + gst_free)


def channels_from_xero() -> list[Channel]:
    splits = profit_and_loss().by_tracking()
    channels = []
    for name in ("North", "South"):
        row = splits[name]
        sales = row.get("Sales - Domestic", 0.0) + row.get("Sales - GST Free", 0.0)
        cogs = abs(row["Cost of Goods Sold"])
        channels.append(
            Channel(
                name=name,
                annual_revenue=sales * 12.0,
                growth_rate=0.0,
                seasonality=[1.0] * 12,
                cogs_pct=cogs / sales,
            )
        )
    return channels


def payroll_frame() -> pd.DataFrame:
    return payroll_forecast(
        ROLES,
        fy_month_range(2027),
        PayrollAssumptions(workers_comp_rate=0.02),
    )


def _monthly_opex() -> list[OpexLine]:
    accounts = profit_and_loss().by_account()
    lines = [
        OpexLine(
            name=account,
            kind="fixed",
            monthly_amount=abs(accounts[account]),
        )
        for account in _NON_PAYROLL_OPEX
    ]
    lines.append(
        OpexLine(
            name="Payroll (statutory)",
            kind="fixed",
            monthly_amount=float(payroll_frame()["total_cost"].iloc[0]),
        )
    )
    return lines


def _working_capital() -> WorkingCapitalConfig:
    accounts = profit_and_loss().by_account()
    bs = balance_sheet().by_account()
    annual_sales = (accounts["Sales - Domestic"] + accounts["Sales - GST Free"]) * 12.0
    annual_cogs = abs(accounts["Cost of Goods Sold"]) * 12.0
    return WorkingCapitalConfig(
        dso_days=bs["Accounts Receivable"] / annual_sales * _DAYS_PER_YEAR,
        dio_days=bs["Inventory"] / annual_cogs * _DAYS_PER_YEAR,
        dpo_days=abs(bs["Accounts Payable"]) / annual_cogs * _DAYS_PER_YEAR,
    )


def entity_config() -> EntityConfig:
    bs = balance_sheet().by_account()
    return EntityConfig(
        name="Harbour Light Pty Ltd",
        start_month="2026-07",
        horizon_months=12,
        tax_rate=0.30,
        channels=channels_from_xero(),
        opex=_monthly_opex(),
        working_capital=_working_capital(),
        opening_balances=OpeningBalances(
            cash=bs["Business Bank Account"],
            ar=bs["Accounts Receivable"],
            ap=abs(bs["Accounts Payable"]),
            inventory=bs["Inventory"],
        ),
    )


def monthly_forecast() -> pd.DataFrame:
    cfg = entity_config()
    frame = cashflow_from_config(cfg)
    frame["leave_provisions"] = payroll_frame()["leave_provisions"]
    revenue, purchases = _gst_inputs()
    frame["gst_accrued"] = monthly_gst(
        revenue, purchases, GstAssumptions(taxable_sales_pct=taxable_sales_pct()),
    )["net_gst"]
    frame["gst_settled"] = 0.0
    for settlement in bas_settlement().itertuples(index=False):
        month = pd.Period(settlement.due_date, freq="M")
        if month in frame.index:
            frame.loc[month, "gst_settled"] += settlement.amount
    frame["gst_cash_impact"] = frame["gst_accrued"] - frame["gst_settled"]
    frame["gst_payable"] = opening_gst() + frame["gst_cash_impact"].cumsum()
    frame["operating_cash_flow"] += frame["leave_provisions"] + frame["gst_cash_impact"]
    frame["free_cash_flow"] = frame["operating_cash_flow"] - frame["capex"]
    frame["change_in_cash"] = frame["free_cash_flow"] - frame["principal"]
    frame["ending_cash"] = cfg.opening_balances.cash + frame["change_in_cash"].cumsum()
    return frame


def opening_gst() -> float:
    accounts = balance_sheet().by_account()
    return -(accounts["GST"] + accounts["GST Clearing"])


def _gst_inputs() -> tuple[pd.Series, pd.Series]:
    months = fy_month_range(2027)
    accounts = profit_and_loss().by_account()
    monthly_sales = accounts["Sales - Domestic"] + accounts["Sales - GST Free"]
    monthly_purchases = abs(accounts["Cost of Goods Sold"]) + sum(
        abs(accounts[name]) for name in _NON_PAYROLL_OPEX
    )
    revenue = pd.Series(monthly_sales, index=months)
    purchases = pd.Series(monthly_purchases, index=months)
    return revenue, purchases


def bas_settlement() -> pd.DataFrame:
    revenue, purchases = _gst_inputs()
    gst = monthly_gst(
        revenue,
        purchases,
        GstAssumptions(
            bas_cycle=BasCycle.QUARTERLY,
            taxable_sales_pct=taxable_sales_pct(),
        ),
    )
    schedule = bas_schedule(gst["net_gst"], GstAssumptions(bas_cycle=BasCycle.QUARTERLY))
    opening = pd.DataFrame([{
        "period_label": "Opening GST", "due_date": OPENING_GST_DUE, "amount": opening_gst(),
    }])
    return pd.concat([opening, schedule], ignore_index=True)


def gst_cash13_flows() -> tuple[list[WeeklyFlow], list[WeeklyFlow]]:
    return gst_weekly_flows(bas_settlement(), window_start=CASH13_START)


def cash13_config() -> Cash13Config:
    monthly = monthly_forecast()
    opening = float(monthly.loc[pd.Period("2026-09", freq="M"), "ending_cash"])
    payroll_cash = float(payroll_frame()["total_cash"].sum()) / 26.0
    gst_receipts, gst_disbursements = gst_cash13_flows()
    accounts = profit_and_loss().by_account()
    weekly_collections = (accounts["Sales - Domestic"] + accounts["Sales - GST Free"]) * 12.0 / 52.0
    weekly_opex = sum(abs(accounts[name]) for name in _NON_PAYROLL_OPEX) * 12.0 / 52.0
    weekly_cogs = abs(accounts["Cost of Goods Sold"]) * 12.0 / 52.0
    gst_rate = GstAssumptions().resolved_rate()
    weekly_collections += accounts["Sales - Domestic"] * gst_rate * 12.0 / 52.0
    weekly_opex *= 1.0 + gst_rate
    weekly_cogs *= 1.0 + gst_rate
    return Cash13Config(
        opening_cash=opening,
        weeks=13,
        receipts=[
            WeeklyFlow(
                name="Collections",
                amount=weekly_collections,
                start_week=1,
                recurrence="weekly",
            ),
            *gst_receipts,
        ],
        disbursements=[
            WeeklyFlow(
                name="Payroll",
                amount=payroll_cash,
                start_week=1,
                recurrence="biweekly",
            ),
            WeeklyFlow(
                name="Supplier payments",
                amount=weekly_cogs,
                start_week=1,
                recurrence="weekly",
            ),
            WeeklyFlow(
                name="Operating overhead",
                amount=weekly_opex,
                start_week=1,
                recurrence="weekly",
            ),
            *gst_disbursements,
        ],
    )


def export_workbook(path: str | Path) -> None:
    """Extend the standard workbook with this example's leave and GST cash bridge."""
    cfg = entity_config()
    model_to_excel(cfg, path)
    wb = load_workbook(path)
    assumptions = wb["Assumptions"]
    assumptions.column_dimensions["A"].width = 30
    assumptions.column_dimensions["B"].width = 20
    payroll = payroll_frame().iloc[0]
    for name, value in {
        "leave_share_of_payroll": float(payroll["leave_provisions"] / payroll["total_cost"]),
        "taxable_sales_share": taxable_sales_pct(),
        "gst_rate": GstAssumptions().resolved_rate(),
        "opening_gst": opening_gst(),
    }.items():
        row = assumptions.max_row + 1
        assumptions.cell(row, 1, name)
        add_named_cell(wb, assumptions, name=name, row=row, col=2, value=value)
    ws = wb["Model"]
    ws.column_dimensions["A"].width = 28
    for column in range(2, cfg.horizon_months + 2):
        ws.column_dimensions[get_column_letter(column)].width = 14
    rows = {ws.cell(row, 1).value: row for row in range(2, ws.max_row + 1)}
    for name in ("leave_provisions", "gst_accrued", "gst_settled", "gst_cash_impact", "gst_payable"):
        rows[name] = ws.max_row + 1
        ws.cell(rows[name], 1, name)
    payroll_index = next(i for i, line in enumerate(cfg.opex, 1) if line.name == "Payroll (statutory)")
    months = fy_month_range(2027)
    schedule = bas_settlement()
    for i, month in enumerate(months):
        col = get_column_letter(i + 2)
        previous = get_column_letter(i + 1)
        payroll_ref = f"{col}{rows[f'opex_{payroll_index}']}"
        settlements = []
        for event in schedule.itertuples(index=False):
            if pd.Period(event.due_date, freq="M") != month:
                continue
            if event.period_label == "Opening GST":
                settlements.append("opening_gst")
            else:
                settlements.extend(
                    f"{get_column_letter(j + 2)}{rows['gst_accrued']}"
                    for j, period in enumerate(months)
                    if str(period.asfreq("Q-JUN")) == event.period_label
                )
        prior_gst = "opening_gst" if i == 0 else f"{previous}{rows['gst_payable']}"
        formulas = {
            "leave_provisions": f"={payroll_ref}*leave_share_of_payroll",
            "gst_accrued": f"=({col}{rows['revenue']}*taxable_sales_share-{col}{rows['cogs']}-{col}{rows['opex']}+{payroll_ref})*gst_rate",
            "gst_settled": "=" + "+".join(settlements or ["0"]),
            "gst_cash_impact": f"={col}{rows['gst_accrued']}-{col}{rows['gst_settled']}",
            "gst_payable": f"={prior_gst}+{col}{rows['gst_cash_impact']}",
            "operating_cash_flow": f"={col}{rows['net_income']}+{col}{rows['da']}+{col}{rows['wc_cash_impact']}+{col}{rows['leave_provisions']}+{col}{rows['gst_cash_impact']}",
        }
        for name, formula in formulas.items():
            ws.cell(rows[name], i + 2, formula).number_format = MONEY_FORMAT
    wb.save(path)
