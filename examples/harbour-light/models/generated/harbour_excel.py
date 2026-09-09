"""Export and verify the Harbour Light leave and GST cash bridge."""
from __future__ import annotations

import argparse
from pathlib import Path
from tempfile import TemporaryDirectory

import harbour_model as hm
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pyfpa.au.calendar import fy_month_range
from pyfpa.au.gst import GstAssumptions
from pyfpa.excel.model_workbook import model_to_excel
from pyfpa.excel.toolkit import MONEY_FORMAT, add_named_cell
from pyfpa.excel.verify import verify_workbook


def _build_workbook(path: str | Path) -> None:
    """Extend the standard workbook with this example's leave and GST cash bridge."""
    cfg = hm.entity_config()
    model_to_excel(cfg, path)
    wb = load_workbook(path)
    assumptions = wb["Assumptions"]
    assumptions.column_dimensions["A"].width = 30
    assumptions.column_dimensions["B"].width = 20
    payroll = hm.payroll_frame().iloc[0]
    for name, value in {
        "leave_share_of_payroll": float(payroll["leave_provisions"] / payroll["total_cost"]),
        "taxable_sales_share": hm.taxable_sales_pct(),
        "gst_rate": GstAssumptions().resolved_rate(),
        "opening_gst": hm.opening_gst(),
    }.items():
        row = assumptions.max_row + 1
        assumptions.cell(row, 1, name)
        add_named_cell(wb, assumptions, name=name, row=row, col=2, value=value)
    ws = wb["Model"]
    ws.column_dimensions["A"].width = 28
    for column in range(2, cfg.horizon_months + 2):
        ws.column_dimensions[get_column_letter(column)].width = 14
    rows = {ws.cell(row, 1).value: row for row in range(2, ws.max_row + 1)}
    for name in ("leave_provisions", "gst_accrued", "gst_settled", "gst_cash_impact", "gst_payable"):
        rows[name] = ws.max_row + 1
        ws.cell(rows[name], 1, name)
    payroll_index = next(i for i, line in enumerate(cfg.opex, 1) if line.name == "Payroll (statutory)")
    months = fy_month_range(2027)
    schedule = hm.bas_settlement()
    for i, month in enumerate(months):
        col = get_column_letter(i + 2)
        previous = get_column_letter(i + 1)
        payroll_ref = f"{col}{rows[f'opex_{payroll_index}']}"
        settlements = []
        for event in schedule.itertuples(index=False):
            if pd.Period(event.due_date, freq="M") != month:
                continue
            if event.period_label == "Opening GST":
                settlements.append("opening_gst")
            else:
                settlements.extend(
                    f"{get_column_letter(j + 2)}{rows['gst_accrued']}"
                    for j, period in enumerate(months)
                    if str(period.asfreq("Q-JUN")) == event.period_label
                )
        prior_gst = "opening_gst" if i == 0 else f"{previous}{rows['gst_payable']}"
        formulas = {
            "leave_provisions": f"={payroll_ref}*leave_share_of_payroll",
            "gst_accrued": f"=({col}{rows['revenue']}*taxable_sales_share-{col}{rows['cogs']}-{col}{rows['opex']}+{payroll_ref})*gst_rate",
            "gst_settled": "=" + "+".join(settlements or ["0"]),
            "gst_cash_impact": f"={col}{rows['gst_accrued']}-{col}{rows['gst_settled']}",
            "gst_payable": f"={prior_gst}+{col}{rows['gst_cash_impact']}",
            "operating_cash_flow": f"={col}{rows['net_income']}+{col}{rows['da']}+{col}{rows['wc_cash_impact']}+{col}{rows['leave_provisions']}+{col}{rows['gst_cash_impact']}",
        }
        for name, formula in formulas.items():
            ws.cell(rows[name], i + 2, formula).number_format = MONEY_FORMAT
    wb.save(path)


def export_workbook(path: str | Path) -> None:
    """Replace the destination only after every forecast line verifies."""
    target = Path(path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=".harbour-", dir=target.parent) as directory:
        candidate = Path(directory) / target.name
        _build_workbook(candidate)
        expected = hm.monthly_forecast()
        report = verify_workbook(candidate, expected)
        if not report.passed or report.lines_checked != len(expected.columns):
            raise ValueError(f"Harbour Light workbook verification failed: {report.failures}")
        candidate.replace(target)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("output/model.xlsx"))
    export_workbook(parser.parse_args().output)
